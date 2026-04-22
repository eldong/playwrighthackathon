"""Generate the hackathon plan DOCX document from the hardcoded agenda content."""

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from pathlib import Path

doc = Document()

style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(11)


def add_styled_table(doc, headers, rows, col_count):
    table = doc.add_table(rows=len(rows) + 1, cols=col_count)
    table.style = 'Light Grid Accent 1'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(headers):
        table.rows[0].cells[j].text = h
        for p in table.rows[0].cells[j].paragraphs:
            for r in p.runs:
                r.bold = True
    for i, row_data in enumerate(rows):
        for j, cell_text in enumerate(row_data):
            table.rows[i + 1].cells[j].text = cell_text
    return table


# --- Title ---
title = doc.add_heading('GitHub Copilot Hackathon', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# --- Objective ---
doc.add_heading('Objective', level=1)
doc.add_paragraph(
    'Experience how GitHub Copilot can accelerate modern software development '
    'by building a simple, end-to-end application within a few hours.'
)
doc.add_paragraph('Participants will use Copilot to:')
for item in [
    'Rapidly scaffold application components (frontend + API)',
    'Implement core functionality using natural language prompts',
    'Apply basic business logic or validation to real-world scenarios',
]:
    doc.add_paragraph(item, style='List Bullet')

doc.add_heading('Theme: Document Validation', level=2)
p = doc.add_paragraph()
run = p.add_run('Build an app that validates documents such as HR forms, invoices, certificates, or ID documents.')
run.italic = True

doc.add_paragraph(
    'The goal is not to build a production-ready system, but to demonstrate how '
    'AI-assisted development can significantly reduce effort, improve productivity, '
    'and enable faster innovation.'
)

# --- Participants ---
doc.add_heading('Participants', level=1)
for item in [
    'Work individually, with optional pairing if preferred',
    'Optional: work in a GitHub repository to collect artifacts after the event',
]:
    doc.add_paragraph(item, style='List Bullet')

# --- Pre-Requisites ---
doc.add_heading('Pre-Requisites (Send 3+ Days Before)', level=1)
doc.add_paragraph('Participants must have the following ready before the hackathon day:')
for item in [
    'VS Code installed',
    'GitHub Copilot extension installed and licensed',
    'Git installed',
    'Runtime of choice installed (Node.js 20+, Python 3.11+, Java 17+, or .NET 8+)',
    'Familiarity with at least one web framework',
]:
    doc.add_paragraph(item, style='List Bullet')

# --- Agenda ---
doc.add_heading('Agenda', level=1)
add_styled_table(doc, ['Phase', 'Duration'], [
    ('Kickoff & Setup', '30 min'),
    ('Build Phase 1 \u2014 Core', '2 hours'),
    ('Build Phase 2 \u2014 Enhance', '1.5 hours'),
    ('Build Phase 3 \u2014 Stretch', '1 hour'),
    ('Demos & Wrap-Up', '1 hour'),
], 2)
doc.add_paragraph()

# Kickoff
doc.add_heading('Kickoff & Setup (30 min)', level=2)
for item in [
    'Welcome and session goals',
    'Live demo: Build a small feature with Copilot in 5 minutes (shows agent mode, inline completions, chat)',
    'Overview of the challenge and example use cases',
    'Verify environments are working',
]:
    doc.add_paragraph(item, style='List Bullet')

# Build Phase 1
doc.add_heading('Build Phase 1 \u2014 Core Flow (2 hours)', level=2)
p = doc.add_paragraph('Goal: ')
run = p.add_run('Achieve a working end-to-end flow.')
run.bold = True
for item in [
    'Create a simple frontend (upload form or input UI)',
    'Build an API endpoint to receive and process data',
    'Implement basic data extraction or parsing logic',
    'Connect frontend \u2192 API \u2192 response displayed to user',
]:
    doc.add_paragraph(item, style='List Bullet')
p = doc.add_paragraph()
run = p.add_run('Tip: ')
run.bold = True
p.add_run('Use a starter template (see below) to skip boilerplate and jump straight to the interesting parts.')

# Build Phase 2
doc.add_heading('Build Phase 2 \u2014 Enhance (1.5 hours)', level=2)
p = doc.add_paragraph('Goal: ')
run = p.add_run('Add business value and polish.')
run.bold = True
for item in [
    'Add validation rules or business logic (required fields, date checks, format validation)',
    'Enhance results display (clear pass/fail indicators, detailed feedback, insights)',
    'Improve user experience (error handling, loading states, styling)',
]:
    doc.add_paragraph(item, style='List Bullet')

# Build Phase 3
doc.add_heading('Build Phase 3 \u2014 Stretch (1 hour)', level=2)
p = doc.add_paragraph('Goal: ')
run = p.add_run('Go beyond the basics.')
run.bold = True
p.add_run(' For teams that have a solid working app.')
for item in [
    'Add unit tests using Copilot-generated test suggestions',
    'Support multiple document types \u2014 add a second validation scenario (e.g., invoices + HR forms) with different rules',
    'Accept PDF uploads \u2014 extract text from uploaded files instead of form input',
    'Build a summary dashboard showing all validation results',
    'Add data persistence (save and retrieve past validations)',
    'Create a reusable Copilot SKILL file \u2014 write a .md skill that captures your validation rules, patterns, or domain knowledge so Copilot can use it in future projects',
    'Add custom instructions (.github/copilot-instructions.md) \u2014 define project-level rules Copilot follows automatically, e.g., consistent error response format or validation patterns',
    'Build a custom agent (.agent.md) \u2014 define a specialized agent with specific instructions, e.g., a \u201cvalidation expert\u201d that knows your document schema and suggests validation patterns',
]:
    doc.add_paragraph(item, style='List Bullet')

# Demos
doc.add_heading('Demos & Wrap-Up (1 hour)', level=2)
for item in [
    'Team presentations (5\u20138 minutes each)',
    'Share what worked, what surprised you, and favorite Copilot moments',
    'Feedback survey',
]:
    doc.add_paragraph(item, style='List Bullet')

# --- Example Use Cases ---
doc.add_heading('Example Use Cases', level=1)
doc.add_paragraph('Pick one to keep scope manageable:')
t = add_styled_table(doc, ['Use Case', 'What It Validates'], [
    ('Invoice Checker', 'Required fields present (vendor, amount, date, PO number), amounts are positive, date is not in the future'),
    ('HR Form Validator', 'Employee name, ID, department filled in; signature present; dates are logical'),
    ('ID Document Verifier', 'Expiry date check, required fields present, format consistency'),
    ('Certificate Validator', 'Issuer information present, valid date range, certificate number format'),
    ('Expense Report Auditor', 'Line items sum to total, receipts attached, within policy limits'),
], 2)
# Bold use case names
for i in range(1, len(t.rows)):
    for p in t.rows[i].cells[0].paragraphs:
        for r in p.runs:
            r.bold = True

# --- Recommended Starter Stacks ---
doc.add_heading('Recommended Starter Stacks', level=1)
doc.add_paragraph('Teams are free to choose their own, but here are four fast-start paths:')
add_styled_table(doc, ['Stack', 'Frontend', 'Backend', 'Good For'], [
    ('JavaScript', 'React or plain HTML', 'Node.js + Express', 'JS/TS-comfortable teams'),
    ('Python', 'HTML + Jinja or React', 'Flask / FastAPI', 'Data/ML-leaning teams'),
    ('Java', 'Thymeleaf or React', 'Spring Boot', 'Java teams'),
    ('.NET', 'Blazor or Razor Pages', 'ASP.NET Core Web API', '.NET teams'),
], 4)

# --- Scope Guidance ---
doc.add_heading('Scope Guidance', level=1)
doc.add_paragraph('To ensure success within the session:')
for item in [
    'Focus on one document type / use case \u2014 don\u2019t try to do everything',
    'Prioritize a working solution over complexity \u2014 a simple app that works beats a complex one that doesn\u2019t',
    'Use Copilot throughout \u2014 the point is the experience, not just the output',
    'Commit early and often \u2014 it\u2019s good practice and helps you track progress',
    'Keep it simple: form-based input is fine; file upload/parsing is a stretch goal',
]:
    doc.add_paragraph(item, style='List Bullet')

# --- What Success Looks Like ---
doc.add_heading('What Success Looks Like', level=1)
for item in [
    'A working application (frontend + API) running locally',
    'Clear demonstration of Copilot-assisted development during the presentation',
    'A practical, business-relevant validation scenario',
    'A creative approach and thoughtful user experience',
    'A team that learned something new about AI-assisted development',
]:
    doc.add_paragraph(item, style='List Bullet')

# Save
output_dir = Path(__file__).parent / 'docs'
output_dir.mkdir(exist_ok=True)
doc.save(output_dir / 'hackathon-plan.docx')
print('Done!')
