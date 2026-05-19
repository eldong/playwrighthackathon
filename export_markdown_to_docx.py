#!/usr/bin/env python3
"""
Convert one or more markdown files to DOCX format.
Usage: python export_markdown_to_docx.py <input.md> [more.md ...]
"""

import re
import sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

update 2 in branch eldong-patch-1


def md_to_docx(markdown_file, output_file=None):
    """Convert markdown file to DOCX document."""
    
    if output_file is None:
        project_root = Path(__file__).parent
        output_dir = project_root / 'docs'
        output_dir.mkdir(exist_ok=True)
        output_file = output_dir / f"{Path(markdown_file).stem}.docx"
    
    doc = Document()
    
    # Read markdown file
    with open(markdown_file, 'r', encoding='utf-8') as f:
        content = f.read()
    
    # Split into lines for processing
    lines = content.split('\n')
    
    i = 0
    while i < len(lines):
        line = lines[i]
        
        # Headings
        if line.startswith('# '):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith('## '):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith('### '):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith('#### '):
            doc.add_heading(line[5:].strip(), level=4)
        
        # Horizontal rule
        elif line.strip() == '---':
            doc.add_paragraph()  # Add spacing
        
        # Tables
        elif line.strip().startswith('|'):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i])
                i += 1
            i -= 1
            
            if len(table_lines) >= 2:
                # Parse table header
                header_row = [cell.strip() for cell in table_lines[0].split('|')[1:-1]]
                
                # Skip separator row (table_lines[1])
                # Parse data rows
                data_rows = []
                for row_line in table_lines[2:]:
                    cells = [cell.strip() for cell in row_line.split('|')[1:-1]]
                    if cells and any(cells):  # Only if row has content
                        data_rows.append(cells)
                
                if data_rows:
                    # Create table
                    table = doc.add_table(rows=1, cols=len(header_row))
                    table.style = 'Light Grid Accent 1'
                    
                    # Add header
                    for j, cell_text in enumerate(header_row):
                        cell = table.rows[0].cells[j]
                        cell.text = cell_text
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                run.bold = True
                    
                    # Add data rows
                    for row_data in data_rows:
                        row_cells = table.add_row().cells
                        for j, cell_text in enumerate(row_data):
                            if j < len(row_cells):
                                row_cells[j].text = cell_text
                    
                    doc.add_paragraph()  # Add spacing
        
        # Bullet lists
        elif line.strip().startswith('- '):
            bullet_text = line.strip()[2:].strip()
            doc.add_paragraph(bullet_text, style='List Bullet')
        
        # Numbered lists
        elif re.match(r'^\d+\.\s', line.strip()):
            match = re.match(r'^(\d+)\.\s(.*)', line.strip())
            if match:
                bullet_text = match.group(2).strip()
                doc.add_paragraph(bullet_text, style='List Number')
        
        # Code blocks
        elif line.strip().startswith('```'):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i])
                i += 1
            
            if code_lines:
                code_text = '\n'.join(code_lines).strip()
                p = doc.add_paragraph(code_text)
                p.style = 'Normal'
                for run in p.runs:
                    run.font.name = 'Courier New'
                    run.font.size = Pt(9)
                doc.add_paragraph()  # Spacing
        
        # Blockquote
        elif line.strip().startswith('> '):
            quote_text = line.strip()[2:].strip()
            p = doc.add_paragraph(quote_text)
            p.paragraph_format.left_indent = Inches(0.5)
            for run in p.runs:
                run.italic = True
        
        # Empty lines (paragraph break)
        elif not line.strip():
            if i > 0 and lines[i-1].strip():  # Only if previous line had content
                pass  # Skip multiple empty lines
        
        # Regular paragraph
        elif line.strip():
            para_text = line.strip()
            doc.add_paragraph(para_text)
        
        i += 1
    
    # Save document
    doc.save(output_file)
    print(f"✓ Created {output_file}")


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print("Usage: python export_markdown_to_docx.py <input.md> [more.md ...]")
        sys.exit(1)

    for input_file in sys.argv[1:]:
        if not Path(input_file).exists():
            print(f"Error: {input_file} not found")
            sys.exit(1)

        md_to_docx(input_file)
